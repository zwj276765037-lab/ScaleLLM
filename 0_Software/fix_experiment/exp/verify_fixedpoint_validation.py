#!/usr/bin/env python3
"""Fail-closed evidence audit for the Llama-3.1-8B format-v3 study.

This verifier checks completeness and provenance, not whether the hardware
quality thresholds pass.  A fully measured failure is a completed validation;
missing tasks, partial PPL, incomplete traces, or a corrupt workbook are not.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TASKS = (
    "piqa",
    "arc_easy",
    "arc_challenge",
    "boolq",
    "hellaswag",
    "winogrande",
)
TARGET_MODEL_NAME = "Llama-3.1-8B-Instruct"
EXPECTED_CASE_METHOD = {
    "v5_attention": "v5_attention_only_controlled_ablation",
    "v5_mlp": "v5_mlp_only_controlled_ablation",
    "v5": "v5_production_fullchain",
}
EXPECTED_SAMPLES = {
    "piqa": 1838,
    "arc_easy": 2376,
    "arc_challenge": 1172,
    "boolq": 3270,
    "hellaswag": 10042,
    "winogrande": 1267,
}
EXPECTED_PRIMARY_METRIC = {
    "piqa": "acc_norm",
    "arc_easy": "acc_norm",
    "arc_challenge": "acc_norm",
    "boolq": "acc",
    "hellaswag": "acc_norm",
    "winogrande": "acc",
    "wikitext2": "token_perplexity",
}
EXPECTED_TRACE_NODES = {
    "attention": {
        "rope_q_local",
        "rope_k_local",
        "kv12_key",
        "kv12_value",
        "qk_score_cumulative",
        "softmax_local",
        "pv_local",
    },
    "mlp": {
        "mlp_gate_projection",
        "mlp_up_projection",
        "mlp_silu_lut_local",
        "mlp_swiglu_round_local",
        "mlp_swiglu_cumulative",
        "mlp_down_projection_local",
        "mlp_down_projection_cumulative",
    },
}
REQUIRED_WIDTH_FINDINGS = {
    "post_rmsnorm_s24_f16",
    "attention_shift_25",
    "attention_qk_pm16_clamp",
    "softmax_probability_u16_counter",
    "mlp_internal_s37_s44_s46",
    "mlp_down_s24_f16",
    "mlp_down_output_abi_chain",
    "mlp_down_s32_validation",
    "artifact_mlp_format_id",
}
REQUIRED_WORKBOOK_SHEETS = {
    "README",
    "Case Summary",
    "ACC",
    "PPL",
    "Trace Coverage",
    "Node Trace",
    "Hardware Audit",
    "Linear Audit",
    "Format Contract",
    "Width Findings",
    "Regression",
    "Provenance",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--baseline-manifest", type=Path, required=True)
    parser.add_argument("--require-complete", action="store_true")
    return parser.parse_args()


def load_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_json(path: Path, value: Any) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    temporary.replace(path)


def main() -> None:
    args = parse_args()
    root = args.root.resolve()
    checks: list[dict[str, Any]] = []

    def check(check_id: str, passed: bool, evidence: Any) -> None:
        checks.append(
            {"check_id": check_id, "pass": bool(passed), "evidence": evidence}
        )

    baseline_path = args.baseline_manifest.resolve()
    checkpoint_identity_path = (
        root.parent
        / "six_model_algorithm_literature_2048"
        / "model_identities"
        / f"{TARGET_MODEL_NAME}.json"
    )
    checkpoint_identity_present = checkpoint_identity_path.is_file()
    check(
        "checkpoint_identity_present",
        checkpoint_identity_present,
        str(checkpoint_identity_path),
    )
    if checkpoint_identity_present:
        checkpoint_identity = load_json(checkpoint_identity_path)
        checkpoint_path = Path(
            str(checkpoint_identity.get("resolved_path") or "")
        ).resolve()
        expected_files = {
            str(row.get("name")): int(row.get("bytes") or -1)
            for row in checkpoint_identity.get("files", [])
        }
        observed_files = (
            {
                path.name: path.stat().st_size
                for path in checkpoint_path.iterdir()
                if path.is_file()
            }
            if checkpoint_path.is_dir()
            else {}
        )
        check(
            "checkpoint_identity_metadata_match",
            checkpoint_identity.get("schema")
            == "scalellm_checkpoint_identity_v1"
            and checkpoint_identity.get("checkpoint_name") == TARGET_MODEL_NAME
            and bool(checkpoint_identity.get("fingerprint_sha256"))
            and int(checkpoint_identity.get("file_count") or 0)
            == len(expected_files)
            and expected_files == observed_files
            and int(checkpoint_identity.get("total_bytes") or 0)
            == sum(observed_files.values()),
            {
                "resolved_path": str(checkpoint_path),
                "fingerprint_sha256": checkpoint_identity.get(
                    "fingerprint_sha256"
                ),
                "expected_file_count": len(expected_files),
                "observed_file_count": len(observed_files),
                "expected_total_bytes": checkpoint_identity.get("total_bytes"),
                "observed_total_bytes": sum(observed_files.values()),
            },
        )
        expected_hashes = {
            str(row.get("name")): str(row.get("sha256") or "")
            for row in checkpoint_identity.get("files", [])
        }
        observed_hashes = {
            name: sha256_file(checkpoint_path / name)
            for name in sorted(expected_hashes)
            if (checkpoint_path / name).is_file()
        }
        hash_mismatches = {
            name: {
                "expected": expected_hashes.get(name),
                "observed": observed_hashes.get(name),
            }
            for name in sorted(set(expected_hashes) | set(observed_hashes))
            if expected_hashes.get(name) != observed_hashes.get(name)
        }
        check(
            "checkpoint_identity_byte_hash_match",
            len(observed_hashes) == len(expected_hashes)
            and not hash_mismatches,
            {
                "files_verified": len(observed_hashes),
                "fingerprint_sha256": checkpoint_identity.get(
                    "fingerprint_sha256"
                ),
                "mismatches": hash_mismatches,
            },
        )
    case_paths = {
        "bf16": baseline_path,
        "v5_attention": root
        / "full_acc_rtl_zero_guard_v3/v5_attention/merged/manifest.json",
        "v5_mlp": root / "full_acc_rtl_zero_guard_v3/v5_mlp/merged/manifest.json",
        "v5": root / "full_acc_rtl_zero_guard_v3/v5/merged/manifest.json",
    }
    manifests: dict[str, dict[str, Any]] = {}
    for case, path in case_paths.items():
        present = path.is_file()
        check(f"{case}_manifest_present", present, str(path))
        if not present:
            continue
        manifest = load_json(path)
        manifests[case] = manifest
        results = manifest.get("results") or []
        by_task = {str(row.get("task")): row for row in results}
        check(
            f"{case}_identity",
            Path(str(manifest.get("model") or "")).name == TARGET_MODEL_NAME
            and manifest.get("case") == case,
            {
                "model": manifest.get("model"),
                "case": manifest.get("case"),
                "expected_model": TARGET_MODEL_NAME,
                "expected_case": case,
            },
        )
        if case != "bf16":
            model_spec = manifest.get("model_spec") or {}
            check(
                f"{case}_model_spec",
                model_spec.get("name") == TARGET_MODEL_NAME
                and model_spec.get("architecture") == "LlamaForCausalLM"
                and int(model_spec.get("layers") or 0) == 32
                and int(model_spec.get("expected_module_count") or 0) == 224,
                model_spec,
            )
        check(
            f"{case}_run_status",
            manifest.get("status") == "complete"
            and manifest.get("hardware_contract_status") == "valid",
            {
                "status": manifest.get("status"),
                "hardware_contract_status": manifest.get(
                    "hardware_contract_status"
                ),
            },
        )
        expected_tasks = set(TASKS) | {"wikitext2"}
        check(
            f"{case}_task_coverage",
            len(results) == len(expected_tasks) and set(by_task) == expected_tasks,
            {
                "rows": len(results),
                "observed": sorted(by_task),
                "expected": sorted(expected_tasks),
            },
        )
        protocol = manifest.get("protocol") or {}
        check(
            f"{case}_hardware128_context",
            protocol.get("protocol") == "hardware_128"
            and protocol.get("case") == case
            and int(protocol.get("context_length") or 0) == 128
            and protocol.get("zero_shot") is True
            and protocol.get("full_split") is True
            and protocol.get("hardware_constrained") is True,
            protocol,
        )
        if case != "bf16":
            method = manifest.get("method") or {}
            common_method_valid = (
                method.get("case") == EXPECTED_CASE_METHOD[case]
                and int(method.get("layers") or 0) == 32
                and int(method.get("max_context") or 0) == 128
                and int(method.get("activation_bits") or 0) == 8
                and method.get("software_reference_not_hardware") is False
            )
            attention_method_valid = case == "v5_mlp" or (
                method.get("attention_format_id") == "0x48384103"
                and int(method.get("attention_weight_bits") or 0) == 8
                and int(method.get("kv_bits") or 0) == 12
            )
            mlp_guard_width = method.get(
                "mlp_group_guard_compute_width",
                method.get("mlp_group_guard_width"),
            )
            gate_up_width = method.get(
                "mlp_gate_up_accumulator_compute_width",
                method.get("mlp_gate_up_accumulator_width"),
            )
            down_width = method.get(
                "mlp_down_accumulator_compute_width",
                method.get("mlp_down_accumulator_width"),
            )
            mlp_method_valid = case == "v5_attention" or (
                method.get("mlp_full_format_id") == "0x48384603"
                and int(method.get("mlp_weight_bits") or 0) == 4
                and method.get("mlp_total_shift_range") == [20, 44]
                and int(mlp_guard_width or 0) == 37
                and int(gate_up_width or 0) == 44
                and int(down_width or 0) == 46
            )
            check(
                f"{case}_format_v3_method",
                common_method_valid
                and attention_method_valid
                and mlp_method_valid,
                method,
            )
            audit = manifest.get("hardware_audit_total") or {}
            check(
                f"{case}_format_v3_numeric_contract",
                int(audit.get("activation_scale_format_error") or 0) == 0
                and int(audit.get("linear_total_shift_format_error") or 0) == 0
                and int(audit.get("linear_group_guard_overflow_s37") or 0) == 0
                and int(audit.get("linear_accumulator_overflow") or 0) == 0,
                {
                    key: audit.get(key)
                    for key in (
                        "activation_scale_format_error",
                        "linear_total_shift_format_error",
                        "linear_group_guard_overflow_s37",
                        "linear_accumulator_overflow",
                    )
                },
            )
        for task in TASKS:
            row = by_task.get(task) or {}
            check(
                f"{case}_{task}_complete",
                row.get("kind") == "acc"
                and row.get("primary_metric") == EXPECTED_PRIMARY_METRIC[task]
                and int(row.get("samples") or 0) == EXPECTED_SAMPLES[task]
                and math.isfinite(float(row.get("value", math.nan)))
                and math.isfinite(float(row.get("stderr", math.nan))),
                row,
            )
        ppl = by_task.get("wikitext2") or {}
        check(
            f"{case}_wikitext2_full",
            ppl.get("kind") == "ppl"
            and ppl.get("primary_metric") == EXPECTED_PRIMARY_METRIC["wikitext2"]
            and ppl.get("full_stream") is True
            and int(ppl.get("evaluated_blocks") or 0) == 2257
            and int(ppl.get("scoring_tokens") or 0) == 286639
            and int(ppl.get("raw_token_count") or 0) == 288936
            and int(ppl.get("block_size") or 0) == 128
            and int(ppl.get("discarded_tail_tokens") or 0) == 40
            and ppl.get("non_overlapping") is True
            and ppl.get("tail_policy") == "drop_incomplete_block"
            and math.isfinite(float(ppl.get("value", math.nan))),
            ppl,
        )
        acc_values = [float(by_task[task]["value"]) for task in TASKS if task in by_task]
        recomputed = sum(acc_values) / 6 if len(acc_values) == 6 else None
        reported = manifest.get("macro_accuracy")
        check(
            f"{case}_macro_recomputed",
            case == "bf16"
            or (
                recomputed is not None
                and reported is not None
                and math.isclose(recomputed, float(reported), rel_tol=0, abs_tol=1e-15)
            ),
            {"reported": reported, "recomputed": recomputed},
        )

    trace_definitions = {
        "attention": root / "node_trace/manifest.json",
        "mlp": root / "node_trace_mlp/manifest.json",
    }
    for domain, path in trace_definitions.items():
        present = path.is_file()
        check(f"{domain}_trace_manifest_present", present, str(path))
        if not present:
            continue
        rows = (load_json(path).get("node_trace") or {}).get("rows") or []
        trace_manifest = load_json(path)
        expected_case = "v5_attention" if domain == "attention" else "v5_mlp"
        trace_protocol = trace_manifest.get("protocol") or {}
        trace_results = trace_manifest.get("results") or []
        trace_ppl = trace_results[0] if len(trace_results) == 1 else {}
        check(
            f"{domain}_trace_provenance",
            Path(str(trace_manifest.get("model") or "")).name
            == TARGET_MODEL_NAME
            and trace_manifest.get("case") == expected_case
            and trace_manifest.get("status") == "complete"
            and trace_manifest.get("hardware_contract_status") == "valid"
            and trace_protocol.get("protocol") == "hardware_128"
            and int(trace_protocol.get("context_length") or 0) == 128
            and trace_ppl.get("task") == "wikitext2"
            and int(trace_ppl.get("evaluated_blocks") or 0) == 1
            and int(trace_ppl.get("scoring_tokens") or 0) == 127,
            {
                "model": trace_manifest.get("model"),
                "case": trace_manifest.get("case"),
                "status": trace_manifest.get("status"),
                "hardware_contract_status": trace_manifest.get(
                    "hardware_contract_status"
                ),
                "protocol": trace_protocol,
                "result": trace_ppl,
            },
        )
        observed = {
            layer: {
                str(row.get("node"))
                for row in rows
                if int(row.get("layer", -1)) == layer
            }
            for layer in (0, 15, 31)
        }
        check(
            f"{domain}_trace_coverage",
            len(rows) == 21
            and all(observed[layer] == EXPECTED_TRACE_NODES[domain] for layer in observed)
            and all(int(row.get("tokens") or 0) == 128 for row in rows),
            {"rows": len(rows), "nodes_by_layer": {str(k): sorted(v) for k, v in observed.items()}},
        )
        finite = all(
            math.isfinite(float(row.get(metric, math.nan)))
            for row in rows
            for metric in ("cosine", "nmse", "mae", "max_abs_error")
        )
        check(f"{domain}_trace_finite", finite, f"rows={len(rows)}")

    contract_path = root / "format_v3_contract_audit.json"
    contract_present = contract_path.is_file()
    check("format_contract_present", contract_present, str(contract_path))
    if contract_present:
        contract = load_json(contract_path)
        core = [row for row in contract.get("checks", []) if row.get("severity") == "core"]
        check(
            "format_contract_core_aligned",
            len(core) == 38 and all(row.get("pass") is True for row in core),
            {"core_checks": len(core), "failed": [row for row in core if not row.get("pass")]},
        )

    width_path = root / "width_findings.csv"
    width_present = width_path.is_file()
    check("width_findings_present", width_present, str(width_path))
    if width_present:
        rows = load_csv(width_path)
        observed = {row.get("finding_id") for row in rows}
        check(
            "width_findings_complete",
            REQUIRED_WIDTH_FINDINGS.issubset(observed),
            {"observed": sorted(observed), "required": sorted(REQUIRED_WIDTH_FINDINGS)},
        )

    report_path = root / "precision_validation_report.md"
    check(
        "markdown_report_present",
        report_path.is_file() and report_path.stat().st_size > 1000,
        str(report_path),
    )
    regression_path = root / "regression_evidence.json"
    regression_present = regression_path.is_file()
    check("regression_evidence_present", regression_present, str(regression_path))
    if regression_present:
        regression = load_json(regression_path)
        check(
            "regression_suites_complete",
            regression.get("status") == "complete_pass"
            and int(regression.get("passed_tests") or 0) == 57
            and all(
                row.get("status") == "pass"
                for row in regression.get("suites", [])
            ),
            {
                "status": regression.get("status"),
                "passed_tests": regression.get("passed_tests"),
                "suites": [
                    {"suite": row.get("suite"), "status": row.get("status")}
                    for row in regression.get("suites", [])
                ],
            },
        )
    workbook_path = root / "ScaleLLM_Llama31_format_v3_validation.xlsx"
    workbook_present = workbook_path.is_file()
    check("workbook_present", workbook_present, str(workbook_path))
    if workbook_present:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        sheetnames = set(workbook.sheetnames)
        workbook.close()
        check(
            "workbook_required_sheets",
            REQUIRED_WORKBOOK_SHEETS.issubset(sheetnames),
            {"observed": sorted(sheetnames), "required": sorted(REQUIRED_WORKBOOK_SHEETS)},
        )
        workbook_manifest_path = root / "validation_workbook_manifest.json"
        manifest_present = workbook_manifest_path.is_file()
        check("workbook_manifest_present", manifest_present, str(workbook_manifest_path))
        if manifest_present:
            workbook_manifest = load_json(workbook_manifest_path)
            check(
                "workbook_sha256",
                workbook_manifest.get("sha256") == sha256_file(workbook_path),
                {
                    "recorded": workbook_manifest.get("sha256"),
                    "actual": sha256_file(workbook_path),
                },
            )

    passed = all(row["pass"] for row in checks)
    quality_outcomes = {
        case: {
            "status": manifest.get("status"),
            "hardware_contract_status": manifest.get("hardware_contract_status"),
            "macro_accuracy": manifest.get("macro_accuracy"),
            "ppl": manifest.get("ppl")
            or (manifest.get("results") or [{}])[-1].get("value"),
        }
        for case, manifest in manifests.items()
    }
    report = {
        "schema": "scalellm_fixedpoint_validation_completion_audit_v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "status": "complete_evidence_verified" if passed else "incomplete_evidence",
        "evidence_complete": passed,
        "quality_outcomes": quality_outcomes,
        "checks_passed": sum(row["pass"] for row in checks),
        "checks_total": len(checks),
        "checks": checks,
    }
    atomic_json(root / "validation_completion_audit.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if args.require_complete and not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
