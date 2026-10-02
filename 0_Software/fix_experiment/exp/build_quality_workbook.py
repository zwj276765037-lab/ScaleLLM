#!/usr/bin/env python3
"""Build the paper-facing ScaleCode ACC/PPL analysis workbook."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


WORKSPACE = Path("/home/user/EDA-202/zhuowenjun")
FIX_ROOT = WORKSPACE / "ScaleLLM" / "0_Software" / "fix_experiment"
RESULT_ROOT = FIX_ROOT / "results" / "six_model_algorithm_literature_2048"
EXTENDED_PPL_ROOT = RESULT_ROOT / "extended_ppl_2048"
EXPERIMENT_MANIFEST = RESULT_ROOT / "experiment_manifest.json"
DATA_MANIFEST = RESULT_ROOT / "data_manifest.json"
MODEL_IDENTITY_ROOT = RESULT_ROOT / "model_identities"
BAPE_DATA_MANIFEST = WORKSPACE / "fpga_llm_project" / "data" / "bape_eval" / "manifest.json"
ARCHIVE_ROOT = (
    WORKSPACE
    / "fpga_llm_project"
    / "results"
    / "All_results"
    / "accuracy_MLP_4bit"
    / "quality_accuracy"
)
ARCHIVE_SUMMARY = (
    WORKSPACE
    / "ScaleLLM"
    / "0_Software"
    / "results"
    / "_summary"
    / "final_scalecode_ppl_wikitext2.csv"
)
LEGACY_COMPLETION_ROOT = (
    FIX_ROOT / "results" / "five_config_legacy_64block_completion"
)
FIVE_CONFIG_ACC_TABLE = (
    WORKSPACE
    / "fpga_llm_project"
    / "results"
    / "0_Final_results"
    / "fix_accuracy"
    / "results"
    / "gyrot_aligned_accuracy"
    / "summary_repaired"
    / "gyrot_aligned_paper_style_table_with_gsm8k_raw_observed.csv"
)
MODELS = (
    "DeepSeek-R1-Distill-Llama-8B",
    "GLM-4-9B-0414",
    "Llama-3.1-8B-Instruct",
    "Llama-3.2-3B-Instruct",
    "Qwen2.5-7B-Instruct",
    "Qwen3-8B",
)
FIVE_CONFIGS = (
    "Full Precision BF16",
    "A8W8 FP32-payload",
    "A8W4 FP32-payload",
    "A8W4 FP16-payload",
    "A8W4 ScaleCode",
)
FORMAL_CASES = (
    "bf16",
    "a8w8_fp32",
    "a8w4_fp32",
    "a8w4_fp16",
    "reference_qdq",
)
FORMAL_CONFIG_DISPLAY = {
    "bf16": "Full precision (BF16)",
    "a8w8_fp32": "A8W8 / FP32 scale",
    "a8w4_fp32": "A8W4 / FP32 scale",
    "a8w4_fp16": "A8W4 / FP16 scale",
    "reference_qdq": "A8W4 / ScaleCode",
}
CONFIG_DISPLAY = {
    "Full Precision BF16": "Full precision (BF16)",
    "A8W8 FP32-payload": "A8W8 / FP32 scale",
    "A8W4 FP32-payload": "A8W4 / FP32 scale",
    "A8W4 FP16-payload": "A8W4 / FP16 scale",
    "A8W4 ScaleCode": "A8W4 / ScaleCode",
}

NAVY = "17365D"
BLUE = "5B9BD5"
LIGHT_BLUE = "DDEBF7"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
YELLOW = "FFD966"
LIGHT_YELLOW = "FFF2CC"
RED = "C00000"
LIGHT_RED = "F4CCCC"
WHITE = "FFFFFF"
GRAY = "E7E6E6"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=RESULT_ROOT / "ScaleLLM_ScaleCode_ACC_PPL_analysis.xlsx",
    )
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def read_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extended_ppl_progress() -> dict:
    tasks = (
        "ptb_ppl",
        "c4_fixed_ppl",
        "lambada_full_ppl",
        "imdb_ppl",
        "ag_news_ppl",
    )
    counts = {task: 0 for task in tasks}
    completed_cells: list[dict] = []
    if EXTENDED_PPL_ROOT.exists():
        for path in EXTENDED_PPL_ROOT.glob("*/*/*_ppl.json"):
            try:
                value = read_json(path)
            except Exception:
                continue
            task = str(value.get("task"))
            if (
                task in counts
                and value.get("status") == "complete"
                and bool(value.get("full_split"))
                and math.isfinite(float(value.get("value")))
            ):
                counts[task] += 1
                completed_cells.append(value)
    runner_state_path = EXTENDED_PPL_ROOT / "runner_state.json"
    runner_state = read_json(runner_state_path) if runner_state_path.exists() else {}
    return {
        "expected_cells": len(MODELS) * len(FORMAL_CASES) * len(tasks),
        "completed_cells": len(completed_cells),
        "task_counts": counts,
        "runner_status": runner_state.get("status", "not_started"),
        "runner_created_at": runner_state.get("created_at"),
        "completed_model_cases": len(list(EXTENDED_PPL_ROOT.glob("*/*/manifest.json"))) if EXTENDED_PPL_ROOT.exists() else 0,
        "root": str(EXTENDED_PPL_ROOT),
    }


def archive_ablation() -> list[dict[str, float | str]]:
    archive = {row["Model"]: row for row in read_csv(ARCHIVE_SUMMARY)}
    rows: list[dict[str, float | str]] = []
    for model in MODELS:
        values: dict[str, float] = {}
        for scale_format in ("fp32", "fp16"):
            source = (
                ARCHIVE_ROOT
                / f"{model}__iah_{scale_format}_scale_mlp_w4"
                / "quality_smoke_by_model.csv"
            )
            candidates = [
                row for row in read_csv(source) if row["task"] == "wikitext2"
            ]
            if len(candidates) != 1:
                raise ValueError(f"expected one WikiText row in {source}")
            values[scale_format] = float(candidates[0]["ppl"])
        w8_source = (
            ARCHIVE_ROOT
            / f"{model}__iah_fp32_scale_mlp_w8"
            / "quality_smoke_by_model.csv"
        )
        if not w8_source.exists():
            w8_source = (
                LEGACY_COMPLETION_ROOT
                / f"{model}__iah_fp32_scale_mlp_w8"
                / "quality_smoke_by_model.csv"
            )
        w8_candidates = [
            row for row in read_csv(w8_source) if row["task"] == "wikitext2"
        ]
        if len(w8_candidates) != 1:
            raise ValueError(f"expected one A8W8 WikiText row in {w8_source}")
        values["w8_fp32"] = float(w8_candidates[0]["ppl"])
        bf16 = float(archive[model]["BF16_PPL"])
        scalecode = float(archive[model]["ScaleCode_PPL"])
        rows.append(
            {
                "model": model,
                "bf16": bf16,
                "w8_fp32": values["w8_fp32"],
                "w4_fp32": values["fp32"],
                "w4_fp16": values["fp16"],
                "w4_scalecode": scalecode,
                "w8_fp32_vs_bf16": 100.0 * (values["w8_fp32"] / bf16 - 1.0),
                "w4_fp32_vs_bf16": 100.0 * (values["fp32"] / bf16 - 1.0),
                "w4_fp16_vs_bf16": 100.0 * (values["fp16"] / bf16 - 1.0),
                "w4_scalecode_vs_bf16": 100.0 * (scalecode / bf16 - 1.0),
                "scalecode_vs_w4_fp32": 100.0 * (scalecode / values["fp32"] - 1.0),
            }
        )
    return rows


def five_config_acc_rows() -> list[dict[str, str]]:
    source_rows = read_csv(FIVE_CONFIG_ACC_TABLE)
    by_key = {(row["Model"], row["Method"]): row for row in source_rows}
    result: list[dict[str, str]] = []
    for model in MODELS:
        for method in FIVE_CONFIGS:
            row = by_key.get((model, method))
            if row is None:
                raise ValueError(f"missing five-config ACC row: {model}/{method}")
            result.append(row)
    if len(result) != len(MODELS) * len(FIVE_CONFIGS):
        raise ValueError("five-config ACC matrix is incomplete")
    return result


def setup_sheet(ws, title: str, subtitle: str, end_col: int = 10) -> None:
    ws.sheet_view.showGridLines = False
    end_letter = get_column_letter(end_col)
    ws.merge_cells(f"A1:{end_letter}1")
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{end_letter}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="44546A")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32


def add_table(ws, start_row: int, end_row: int, end_col: int, name: str) -> None:
    reference = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def style_header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    ws.row_dimensions[row].height = 34


def set_widths(ws, widths: dict[int, float]) -> None:
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = width


def formal_five_config_sheet(wb: Workbook, summary: list[dict[str, str]]) -> None:
    ws = wb.active
    ws.title = "五配置总表"
    setup_sheet(
        ws,
        "六模型 × 五配置正式 ACC/PPL 总表",
        "统一正式口径：lm-eval 0.4.11 六任务 full split；WikiText-2 full test、2048-token non-overlap、tail-drop、token-weighted NLL。每个模型固定五行。",
        end_col=12,
    )
    headers = [
        "模型",
        "配置",
        "Macro ACC",
        "Δ ACC vs BF16 (pp)",
        "WikiText-2 PPL",
        "Δ PPL vs BF16",
        "PPL相对变化 vs BF16",
        "PPL blocks",
        "PPL scoring tokens",
        "Case",
        "状态",
        "Checkpoint SHA-256",
    ]
    ws.append([])
    ws.append(headers)
    by_key = {(row["model"], row["case"]): row for row in summary}
    expected = {(model, case) for model in MODELS for case in FORMAL_CASES}
    missing = sorted(expected - set(by_key))
    if missing:
        raise ValueError(f"formal five-config result matrix is incomplete: {missing}")
    for model in MODELS:
        baseline = by_key[(model, "bf16")]
        baseline_acc = float(baseline["macro_accuracy"])
        baseline_ppl = float(baseline["ppl"])
        for case in FORMAL_CASES:
            row = by_key[(model, case)]
            acc = float(row["macro_accuracy"])
            ppl = float(row["ppl"])
            ws.append(
                [
                    model,
                    FORMAL_CONFIG_DISPLAY[case],
                    acc,
                    100.0 * (acc - baseline_acc),
                    ppl,
                    ppl - baseline_ppl,
                    ppl / baseline_ppl - 1.0,
                    int(row["ppl_blocks"]),
                    int(row["ppl_scoring_tokens"]),
                    case,
                    row["status"],
                    row["model_fingerprint_sha256"],
                ]
            )
    header_row = 4
    end_row = header_row + len(MODELS) * len(FORMAL_CASES)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "FormalFiveConfigSummary")
    for row_index in range(header_row + 1, end_row + 1):
        ws.cell(row_index, 3).number_format = "0.0000%"
        ws.cell(row_index, 4).number_format = "+0.0000;-0.0000;0.0000"
        ws.cell(row_index, 5).number_format = "0.000000"
        ws.cell(row_index, 6).number_format = "+0.000000;-0.000000;0.000000"
        ws.cell(row_index, 7).number_format = "+0.0000%;-0.0000%;0.0000%"
        ws.cell(row_index, 8).number_format = "#,##0"
        ws.cell(row_index, 9).number_format = "#,##0"
        if ws.cell(row_index, 2).value == "A8W4 / ScaleCode":
            for col in range(1, len(headers) + 1):
                ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.conditional_formatting.add(
        f"D{header_row + 1}:D{end_row}",
        ColorScaleRule(
            start_type="min", start_color=LIGHT_RED,
            mid_type="num", mid_value=0, mid_color=LIGHT_YELLOW,
            end_type="max", end_color=LIGHT_GREEN,
        ),
    )
    ws.conditional_formatting.add(
        f"G{header_row + 1}:G{end_row}",
        ColorScaleRule(
            start_type="min", start_color=LIGHT_GREEN,
            mid_type="num", mid_value=0, mid_color=LIGHT_YELLOW,
            end_type="max", end_color=LIGHT_RED,
        ),
    )
    ws.freeze_panes = "C5"
    set_widths(
        ws,
        {1: 36, 2: 25, 3: 15, 4: 22, 5: 19, 6: 20, 7: 24, 8: 13, 9: 21, 10: 20, 11: 12, 12: 68},
    )


def main_sheet(wb: Workbook, summary: list[dict[str, str]], deltas: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("ScaleCode结论")
    setup_sheet(
        ws,
        "ScaleLLM ScaleCode 六模型 ACC/PPL 总览",
        "正式口径：BF16 matched baseline vs Hadamard32 + Group32 A8 + Attention W8 + MLP W4 + 8-bit ScaleCode；WikiText-2 full test，2048-token non-overlap。",
    )
    delta_by_model = {row["model"]: row for row in deltas}
    rows_by_key = {(row["model"], row["case"]): row for row in summary}
    headers = [
        "模型",
        "BF16 Macro ACC",
        "ScaleCode Macro ACC",
        "Δ ACC (pp)",
        "BF16 PPL",
        "ScaleCode PPL",
        "Δ PPL",
        "PPL相对变化",
        "ΔNLL (nat/token)",
        "判断",
    ]
    header_row = 4
    ws.append([])
    ws.append(headers)
    for model in MODELS:
        baseline = rows_by_key[(model, "bf16")]
        quantized = rows_by_key[(model, "reference_qdq")]
        delta = delta_by_model[model]
        bf16_ppl = float(baseline["ppl"])
        qdq_ppl = float(quantized["ppl"])
        ppl_relative = float(delta["ppl_relative_change_percent"])
        acc_delta = float(delta["macro_accuracy_delta_pp"])
        if ppl_relative > 8.0:
            judgement = "PPL敏感，不能称无损"
        elif acc_delta < -1.0:
            judgement = "ACC需关注"
        else:
            judgement = "退化相对温和"
        ws.append(
            [
                model,
                float(baseline["macro_accuracy"]),
                float(quantized["macro_accuracy"]),
                acc_delta,
                bf16_ppl,
                qdq_ppl,
                qdq_ppl - bf16_ppl,
                ppl_relative / 100.0,
                math.log(qdq_ppl / bf16_ppl),
                judgement,
            ]
        )
    end_row = header_row + len(MODELS)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "MainResults")
    for row in range(header_row + 1, end_row + 1):
        for col in (2, 3):
            ws.cell(row, col).number_format = "0.0000%"
        ws.cell(row, 4).number_format = "0.0000"
        for col in (5, 6, 7):
            ws.cell(row, col).number_format = "0.00000"
        ws.cell(row, 8).number_format = "+0.0000%;-0.0000%;0.0000%"
        ws.cell(row, 9).number_format = "0.000000"
    ws.conditional_formatting.add(
        f"D{header_row + 1}:D{end_row}",
        ColorScaleRule(
            start_type="min", start_color=LIGHT_RED,
            mid_type="percentile", mid_value=50, mid_color=LIGHT_YELLOW,
            end_type="max", end_color=LIGHT_GREEN,
        ),
    )
    ws.conditional_formatting.add(
        f"H{header_row + 1}:H{end_row}",
        ColorScaleRule(
            start_type="min", start_color=LIGHT_GREEN,
            mid_type="percentile", mid_value=50, mid_color=LIGHT_YELLOW,
            end_type="max", end_color=LIGHT_RED,
        ),
    )
    ws.freeze_panes = "B5"
    set_widths(ws, {1: 36, 2: 18, 3: 21, 4: 14, 5: 14, 6: 17, 7: 13, 8: 17, 9: 19, 10: 24})

    chart_acc = BarChart()
    chart_acc.type = "bar"
    chart_acc.style = 10
    chart_acc.title = "ScaleCode 相对 BF16：Macro ACC 变化"
    chart_acc.y_axis.title = "模型"
    chart_acc.x_axis.title = "百分点（pp）"
    chart_acc.add_data(Reference(ws, min_col=4, min_row=header_row, max_row=end_row), titles_from_data=True)
    chart_acc.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=end_row))
    chart_acc.height = 7.5
    chart_acc.width = 14
    ws.add_chart(chart_acc, "L3")

    chart_ppl = BarChart()
    chart_ppl.type = "bar"
    chart_ppl.style = 10
    chart_ppl.title = "ScaleCode 相对 BF16：PPL 增幅"
    chart_ppl.y_axis.title = "模型"
    chart_ppl.x_axis.title = "相对变化"
    chart_ppl.add_data(Reference(ws, min_col=8, min_row=header_row, max_row=end_row), titles_from_data=True)
    chart_ppl.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=end_row))
    chart_ppl.height = 7.5
    chart_ppl.width = 14
    ws.add_chart(chart_ppl, "L19")

    note_row = end_row + 3
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    ws.cell(note_row, 1, "结论：六模型平均 Macro ACC 变化 -1.0600 pp，平均相对 PPL 增加 7.2975%；DeepSeek 与 Llama-3.2-3B 的 PPL 增幅超过 13%，不能笼统声称无损。")
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
    ws.cell(note_row, 1).font = Font(bold=True, color="7F6000")
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[note_row].height = 36


def task_sheet(wb: Workbook, summary: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("五配置逐任务")
    setup_sheet(
        ws,
        "六模型 × 五配置逐任务完整 ACC",
        "每个模型固定五行；PIQA/ARC/HellaSwag 使用 acc_norm，BoolQ/WinoGrande 使用 acc；所有任务均为 lm-eval 0.4.11 full split。",
        end_col=12,
    )
    headers = [
        "模型", "配置", "PIQA", "ARC-Easy", "ARC-Challenge", "BoolQ",
        "HellaSwag", "WinoGrande", "Macro ACC", "WikiText-2 PPL", "PPL scoring tokens", "状态",
    ]
    ws.append([])
    ws.append(headers)
    row_lookup = {(row["model"], row["case"]): row for row in summary}
    for model in MODELS:
        for case in FORMAL_CASES:
            row = row_lookup[(model, case)]
            ws.append(
                [
                    model,
                    FORMAL_CONFIG_DISPLAY[case],
                    *[float(row[key]) for key in ("piqa", "arc_easy", "arc_challenge", "boolq", "hellaswag", "winogrande")],
                    float(row["macro_accuracy"]),
                    float(row["ppl"]),
                    int(row["ppl_scoring_tokens"]),
                    row["status"],
                ]
            )
    header_row = 4
    end_row = header_row + len(FORMAL_CASES) * len(MODELS)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "TaskResults")
    for row in range(header_row + 1, end_row + 1):
        for col in range(3, 10):
            ws.cell(row, col).number_format = "0.0000%"
        ws.cell(row, 10).number_format = "0.00000"
        ws.cell(row, 11).number_format = "#,##0"
        if ws.cell(row, 2).value == "A8W4 / ScaleCode":
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.freeze_panes = "C5"
    set_widths(ws, {1: 36, 2: 25, 3: 13, 4: 14, 5: 18, 6: 13, 7: 15, 8: 17, 9: 16, 10: 18, 11: 20, 12: 12})


def ablation_sheet(wb: Workbook, rows: list[dict[str, float | str]]) -> None:
    ws = wb.create_sheet("五配置PPL")
    setup_sheet(
        ws,
        "五配置 WikiText-2 PPL 消融",
        "归档五配置口径：WikiText-2 固定 64×1024-token blocks（65,472 scoring tokens）；用于配置间 matched 对比，不与 2048-token 正式主表 PPL 混算。",
    )
    headers = [
        "模型",
        "Full precision (BF16)",
        "A8W8 / FP32 scale",
        "A8W4 / FP32 scale",
        "A8W4 / FP16 scale",
        "A8W4 / ScaleCode",
        "A8W8 FP32 vs BF16",
        "A8W4 FP32 vs BF16",
        "A8W4 FP16 vs BF16",
        "A8W4 ScaleCode vs BF16",
        "ScaleCode vs A8W4 FP32",
    ]
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                row["model"],
                row["bf16"],
                row["w8_fp32"],
                row["w4_fp32"],
                row["w4_fp16"],
                row["w4_scalecode"],
                float(row["w8_fp32_vs_bf16"]) / 100.0,
                float(row["w4_fp32_vs_bf16"]) / 100.0,
                float(row["w4_fp16_vs_bf16"]) / 100.0,
                float(row["w4_scalecode_vs_bf16"]) / 100.0,
                float(row["scalecode_vs_w4_fp32"]) / 100.0,
            ]
        )
    header_row = 4
    end_row = header_row + len(rows)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "PPLAblation")
    for row in range(header_row + 1, end_row + 1):
        for col in range(2, 7):
            ws.cell(row, col).number_format = "0.000000"
        for col in range(7, 12):
            ws.cell(row, col).number_format = "+0.0000%;-0.0000%;0.0000%"
    ws.conditional_formatting.add(
        f"K{header_row + 1}:K{end_row}",
        CellIsRule(operator="between", formula=["-0.002", "0.002"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
    )
    ws.freeze_panes = "B5"
    set_widths(ws, {1: 36, 2: 22, 3: 21, 4: 21, 5: 21, 6: 21, 7: 22, 8: 22, 9: 22, 10: 25, 11: 27})

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "相对 BF16 的 PPL 变化：四种量化配置"
    chart.y_axis.title = "相对增幅"
    chart.x_axis.title = "模型"
    chart.add_data(Reference(ws, min_col=7, max_col=10, min_row=header_row, max_row=end_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=end_row))
    chart.height = 8
    chart.width = 16
    ws.add_chart(chart, "L3")


def five_config_acc_sheet(wb: Workbook, rows: list[dict[str, str]]) -> None:
    ws = wb.create_sheet("论文六任务ACC", 0)
    setup_sheet(
        ws,
        "六模型 × 五配置论文六任务 ACC",
        "按论文表列序：PIQA、ARC-Challenge、OpenBookQA、HellaSwag、WinoGrande 为完整验证集；GSM8K 为 first-300 生成式 exact match。每个模型固定五行。",
        end_col=13,
    )
    headers = [
        "模型",
        "配置",
        "PIQA",
        "ARC-Challenge",
        "OpenBookQA",
        "HellaSwag",
        "WinoGrande",
        "GSM8K",
        "六任务 Avg.",
        "Δ vs BF16 (pp)",
        "Case",
        "GSM8K状态",
        "原始结果文件",
    ]
    ws.append([])
    ws.append(headers)
    bf16_avg = {
        row["Model"]: float(row["Avg."])
        for row in rows
        if row["Method"] == "Full Precision BF16"
    }
    for row in rows:
        avg = float(row["Avg."])
        ws.append(
            [
                row["Model"],
                CONFIG_DISPLAY[row["Method"]],
                float(row["PIQA"]) / 100.0,
                float(row["ARC-c"]) / 100.0,
                float(row["OpenBookQA"]) / 100.0,
                float(row["HellaS."]) / 100.0,
                float(row["WinoG."]) / 100.0,
                float(row["GSM8K"]) / 100.0,
                avg / 100.0,
                avg - bf16_avg[row["Model"]],
                row["Case"],
                row["GSM8K Status"],
                row["Source"],
            ]
        )
    header_row = 4
    end_row = header_row + len(rows)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "FiveConfigACC")
    for row_index in range(header_row + 1, end_row + 1):
        for col in range(3, 10):
            ws.cell(row_index, col).number_format = "0.00%"
        ws.cell(row_index, 10).number_format = "+0.00;-0.00;0.00"
        if ws.cell(row_index, 2).value == "A8W4 / ScaleCode":
            for col in range(1, len(headers) + 1):
                ws.cell(row_index, col).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.conditional_formatting.add(
        f"J{header_row + 1}:J{end_row}",
        ColorScaleRule(
            start_type="min", start_color=LIGHT_RED,
            mid_type="num", mid_value=0, mid_color=LIGHT_YELLOW,
            end_type="max", end_color=LIGHT_GREEN,
        ),
    )
    ws.freeze_panes = "C5"
    set_widths(
        ws,
        {1: 36, 2: 25, 3: 12, 4: 18, 5: 17, 6: 14, 7: 16, 8: 13, 9: 16, 10: 18, 11: 24, 12: 34, 13: 95},
    )


def section_title(ws, row: int, title: str, end_col: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    cell = ws.cell(row, 1, title)
    cell.font = Font(size=12, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def protocol_sheet(wb: Workbook, progress: dict) -> None:
    ws = wb.create_sheet("评测口径")
    setup_sheet(
        ws,
        "ScaleLLM ACC/PPL 冻结评测口径",
        "本页是结果可复现约束。不同 split、prompt、metric、context、token 拼接或量化 case 的结果不得直接混表。",
        end_col=12,
    )
    experiment = read_json(EXPERIMENT_MANIFEST)
    data_rows = read_json(DATA_MANIFEST)["datasets"]
    data_by_task = {row["task"]: row for row in data_rows}
    source_manifest = read_json(Path(experiment["source_manifests"][0]))
    versions = source_manifest["versions"]

    row = 4
    section_title(ws, row, "A. 全局冻结参数", 12)
    row += 1
    general_headers = ["参数", "固定值", "适用范围", "解释/禁止事项"]
    for col, value in enumerate(general_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(general_headers))
    general = [
        ("实验角色", "软件算法 QDQ 质量评测", "全部结果", "不是 FPGA/S24/定点累加器 bit-exact 仿真"),
        ("协议 ID", "literature_2048", "正式 ACC/PPL", "不得与 hardware_128 的绝对值混比"),
        ("软件栈", f"lm-eval {versions['lm_eval']}；PyTorch {versions['torch']}；Transformers {versions['transformers']}；Datasets {versions['datasets']}", "正式 ACC/PPL", "升级版本必须重新跑 BF16 matched baseline"),
        ("模型数 × 配置数", "6 × 5 = 30 model-case", "正式矩阵", "每模型固定五配置，不使用插值或历史值补空"),
        ("随机种子", "20260913", "ACC evaluator", "random/numpy/torch/few-shot seed 固定"),
        ("ACC shot", "zero-shot / num_fewshot=0", "六项 ACC", "不应用 chat template，不添加 system prompt"),
        ("ACC 上下文", "2048 tokens", "六项 ACC", "lm-eval causal continuation likelihood；truncation 统计写入 manifest"),
        ("模型执行", "BF16 compute；eager attention；use_cache=False", "五配置", "同模型五配置共用 checkpoint/tokenizer/prompt"),
        ("ACC 汇总", "六任务等权 Macro ACC", "五配置总表", "任务级宏平均，不按样本数做 micro average"),
        ("PPL 上下文/步长", "block=2048；stride=2048；non-overlap", "全部连续语料 PPL", "仅完整 block；不使用滑窗重叠"),
        ("PPL 公式", "PPL = exp(sum(NLL) / scoring_tokens)", "全部 PPL", "禁止先算每 block PPL 再算术平均"),
        ("PPL token 规则", "add_special_tokens=False；每 block 首 token 只作条件", "全部 PPL", "不足 2048 的 tail 丢弃并记录"),
        ("边界 tokenizer 规则", "fixed_context_independent_continuation_fallback_v1", "非 prefix-stable tokenizer", "只在联合编码不保持 context 前缀时独立编码 continuation；BF16/量化一致"),
        ("硬件排除项", "S24、整数 Hadamard、定点 ACC/KV/Softmax/PV 均不包含", "全部结果", "本文软件表只能证明算法 QDQ 精度"),
    ]
    for values in general:
        row += 1
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        for col in range(1, 5):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 34

    row += 2
    section_title(ws, row, "B. 五种配置的唯一含义", 12)
    row += 1
    config_headers = ["显示名", "Case", "Hadamard", "激活", "Attention 权重", "MLP 权重", "Scale payload", "Compute", "说明"]
    for col, value in enumerate(config_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(config_headers))
    configs = [
        ("Full precision (BF16)", "bf16", "无", "BF16", "BF16", "BF16", "不适用", "BF16", "matched baseline；不插入量化/Hadamard"),
        ("A8W8 / FP32 scale", "a8w8_fp32", "normalized H32", "A8 Group32", "W8", "W8", "float32 scale", "BF16", "全 Linear 算法 QDQ；lm_head/embedding/norm 排除"),
        ("A8W4 / FP32 scale", "a8w4_fp32", "normalized H32", "A8 Group32", "W8 Q/K/V/O", "W4 Gate/Up/Down", "float32 scale", "BF16", "GLM gate_up_proj 按融合 MLP W4"),
        ("A8W4 / FP16 scale", "a8w4_fp16", "normalized H32", "A8 Group32", "W8 Q/K/V/O", "W4 Gate/Up/Down", "float16 scale", "BF16", "仅 scale payload 相对 FP32 改为 FP16"),
        ("A8W4 / ScaleCode", "reference_qdq", "normalized H32", "A8 Group32", "W8 Q/K/V/O", "W4 Gate/Up/Down", "8-bit code + uint8 shift_bits", "BF16", "activation 与 folded-weight scale 均编码；算法参考，不是硬件定点"),
    ]
    for values in configs:
        row += 1
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 48

    row += 2
    section_title(ws, row, "C. ACC 数据、split、样本数与主指标", 12)
    row += 1
    acc_headers = ["Task", "Split", "样本数", "主指标", "选择规则", "Shot", "Source revision", "SHA-256", "本地文件"]
    for col, value in enumerate(acc_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(acc_headers))
    choice_rule = {
        "piqa": "acc_norm：choice 总 log-likelihood / 字符长度",
        "arc_easy": "acc_norm：choice 总 log-likelihood / 字符长度",
        "arc_challenge": "acc_norm：choice 总 log-likelihood / 字符长度",
        "boolq": "acc：choice 总 log-likelihood 最大",
        "hellaswag": "acc_norm：choice 总 log-likelihood / 字符长度",
        "winogrande": "acc：填空后共同 suffix 条件似然最大",
    }
    task_display = {
        "piqa": "PIQA",
        "arc_easy": "ARC-Easy",
        "arc_challenge": "ARC-Challenge",
        "boolq": "BoolQ",
        "hellaswag": "HellaSwag",
        "winogrande": "WinoGrande",
    }
    for task in task_display:
        item = data_by_task[task]
        row += 1
        values = [
            task_display[task], item["split"], item["rows"], item["primary_metric"],
            choice_rule[task], "0-shot", item["source_revision"], item["sha256"], item["path"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 48

    row += 2
    section_title(ws, row, "D. PPL 数据处理与当前覆盖", 12)
    row += 1
    ppl_headers = ["数据集", "固定范围", "行/文档数", "文本拼接", "Tokenizer", "Block", "Stride", "Tail", "统计", "完成度", "SHA-256", "本地文件"]
    for col, value in enumerate(ppl_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(ppl_headers))
    bape_assets = {asset["id"]: asset for asset in read_json(BAPE_DATA_MANIFEST)["assets"]}
    wikitext = data_by_task["wikitext2"]
    ppl_rows = [
        ("WikiText-2 raw-v1", "test/full", 4358, "去空行后以 \\n\\n 拼接", "checkpoint 原生；无 special tokens", 2048, 2048, "drop incomplete", "token-weighted NLL", "30/30 完成", wikitext["sha256"], wikitext["path"]),
    ]
    extended_specs = [
        ("PTB legacy text-only", "ptb_test", "ptb_ppl"),
        ("C4-en fixed 10k", "c4_official_validation_first10000", "c4_fixed_ppl"),
        ("LAMBADA full-text", "lambada_test", "lambada_full_ppl"),
        ("IMDB", "imdb_test", "imdb_ppl"),
        ("AG News", "ag_news_test", "ag_news_ppl"),
    ]
    for display, asset_id, task in extended_specs:
        asset = bape_assets[asset_id]
        ppl_rows.append(
            (
                display, asset["split"], asset["rows"], "文档按固定顺序、separator=none",
                "checkpoint 原生；无 special tokens", 2048, 2048, "drop incomplete",
                "token-weighted NLL；LAMBADA 为全文 PPL", f"{progress['task_counts'][task]}/30 完成",
                asset["sha256"], str(WORKSPACE / "fpga_llm_project" / asset["path"]),
            )
        )
    for values in ppl_rows:
        row += 1
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 54

    ws.freeze_panes = "A5"
    set_widths(
        ws,
        {1: 27, 2: 31, 3: 14, 4: 35, 5: 31, 6: 12, 7: 12, 8: 19, 9: 29, 10: 17, 11: 68, 12: 88},
    )


def audit_sheet(wb: Workbook, progress: dict) -> None:
    ws = wb.create_sheet("实验审计")
    setup_sheet(
        ws,
        "ScaleLLM 结果来源与完整性审计",
        "所有结果均可追溯到 checkpoint fingerprint、数据 SHA-256、代码 aggregate SHA-256 和原始 manifest；本页生成时刻之后的运行进度不会自动回填。",
        end_col=10,
    )
    experiment = read_json(EXPERIMENT_MANIFEST)
    source_manifest = read_json(Path(experiment["source_manifests"][0]))
    versions = source_manifest["versions"]
    row = 4
    section_title(ws, row, "A. 完整性与运行状态快照", 10)
    row += 1
    headers = ["项目", "值", "状态/解释", "证据文件"]
    for col, value in enumerate(headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(headers))
    audit_rows = [
        ("工作簿生成时间", datetime.now().astimezone().isoformat(timespec="seconds"), "静态快照", str(RESULT_ROOT)),
        ("正式协议", experiment["protocol"], experiment["status"], str(EXPERIMENT_MANIFEST)),
        ("正式 model-case", f"{len(experiment['completed_model_cases'])}/30", f"missing={len(experiment['missing'])}；validation_errors={len(experiment['validation_errors'])}", str(EXPERIMENT_MANIFEST)),
        ("正式 ACC 单元", "180/180", "6 models × 5 configs × 6 tasks", str(RESULT_ROOT / "six_model_summary.csv")),
        ("正式 WikiText-2 PPL", "30/30", "6 models × 5 configs", str(RESULT_ROOT / "six_model_summary.csv")),
        ("扩展 PPL 单元", f"{progress['completed_cells']}/{progress['expected_cells']}", f"runner={progress['runner_status']}；完整 model-case={progress['completed_model_cases']}/30", progress["root"]),
        ("代码 aggregate SHA-256", experiment["code_identity"]["aggregate_sha256"], "冻结正式评测实现", str(EXPERIMENT_MANIFEST)),
        ("旧历史数据", "未写入本工作簿", "禁止用旧 gyrot/GSM8K first-300 填补新实验", "见阅读说明"),
        ("硬件 bit-exact", "不包含", "S24/定点 Hadamard/ACC/KV/Softmax/PV 不在本表", "软件算法边界"),
    ]
    for values in audit_rows:
        row += 1
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 40

    row += 2
    section_title(ws, row, "B. 软件环境", 10)
    row += 1
    for col, value in enumerate(["组件", "固定版本", "说明"], 1):
        ws.cell(row, col, value)
    style_header(ws, row, 3)
    for component in ("python", "torch", "transformers", "datasets", "lm_eval", "accelerate", "triton"):
        row += 1
        ws.cell(row, 1, component)
        ws.cell(row, 2, versions[component])
        ws.cell(row, 3, "正式结果 manifest 记录值")

    row += 2
    section_title(ws, row, "C. 六个 checkpoint 身份", 10)
    row += 1
    model_headers = ["模型", "model_type", "架构", "层数", "目标模块数", "文件数", "总大小(GiB)", "Checkpoint fingerprint SHA-256", "Tokenizer SHA-256", "本地路径"]
    for col, value in enumerate(model_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(model_headers))
    for model in MODELS:
        identity = read_json(MODEL_IDENTITY_ROOT / f"{model}.json")
        spec = identity["model_spec"]
        tokenizer_files = [item for item in identity["files"] if item["name"] == "tokenizer.json"]
        tokenizer_sha = tokenizer_files[0]["sha256"] if tokenizer_files else "not present"
        row += 1
        values = [
            model, spec["model_type"], spec["architecture"], spec["layers"],
            spec["expected_module_count"], identity["file_count"], identity["total_bytes"] / (1024 ** 3),
            identity["fingerprint_sha256"], tokenizer_sha, identity["resolved_path"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row, 7).number_format = "0.000"
        ws.row_dimensions[row].height = 52

    row += 2
    section_title(ws, row, "D. Tokenizer 边界兼容审计", 10)
    row += 1
    boundary_headers = ["模型", "总 requests", "fallback requests", "fallback 比例", "修复后空 continuation", "策略"]
    for col, value in enumerate(boundary_headers, 1):
        ws.cell(row, col, value)
    style_header(ws, row, len(boundary_headers))
    for item in experiment["tokenizer_boundary_audits"]:
        row += 1
        total = int(item["total_requests"])
        fallback = int(item["fallback_requests"])
        values = [
            item["checkpoint_name"], total, fallback, fallback / total if total else 0.0,
            item["empty_continuations_after_fix"], item["policy"],
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 4).number_format = "0.0000%"

    row += 2
    section_title(ws, row, "E. 关键结果文件", 10)
    row += 1
    for col, value in enumerate(["文件", "字节数", "修改时间", "SHA-256"], 1):
        ws.cell(row, col, value)
    style_header(ws, row, 4)
    evidence_files = [
        EXPERIMENT_MANIFEST,
        DATA_MANIFEST,
        RESULT_ROOT / "six_model_summary.csv",
        RESULT_ROOT / "six_model_delta.csv",
        RESULT_ROOT / "five_config_delta.csv",
    ]
    for path in evidence_files:
        row += 1
        stat = path.stat()
        values = [str(path), stat.st_size, datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"), sha256_file(path)]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A5"
    set_widths(ws, {1: 40, 2: 22, 3: 30, 4: 24, 5: 18, 6: 18, 7: 18, 8: 70, 9: 70, 10: 90})


def dataset_sheet(wb: Workbook, progress: dict) -> None:
    ws = wb.create_sheet("数据集覆盖")
    setup_sheet(
        ws,
        "本地完整评测集与六模型五配置覆盖状态",
        "本页列出 fpga_llm_project/data 中可用于论文评测的全部主/补充测试集；“数据已就绪”不等于已经完成六模型×五配置实测。",
    )
    datasets = [
        {
            "name": "WikiText-2 raw-v1",
            "role": "PPL 主任务",
            "split": "test/full",
            "rows": 4358,
            "path": WORKSPACE / "fpga_llm_project/data/wikitext-2-raw-v1/test-00000-of-00001.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "非空行以双换行拼接；原生 tokenizer；2048 non-overlap；tail 丢弃；token-weighted NLL",
            "priority": "正式结果",
        },
        {
            "name": "C4-en official fixed subset",
            "role": "PPL 跨域主补充",
            "split": "validation shard0 rows 0..9999",
            "rows": 10000,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/c4_official/validation_first10000.jsonl",
            "status": f"六模型×五配置运行中：{progress['task_counts']['c4_fixed_ppl']}/30 完成",
            "protocol": "固定 source_index；无分隔拼接；原生 tokenizer；2048 non-overlap；tail 丢弃",
            "priority": "运行中",
        },
        {
            "name": "PTB text-only",
            "role": "补充 PPL",
            "split": "test/full",
            "rows": 3761,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/ptb_text_only/ptb.test.txt",
            "status": f"六模型×五配置运行中：{progress['task_counts']['ptb_ppl']}/30 完成",
            "protocol": "保留预处理 <unk>/<N>；无分隔拼接；原生 tokenizer；2048 non-overlap；tail 丢弃",
            "priority": "运行中",
        },
        {
            "name": "LAMBADA OpenAI",
            "role": "补充全文 PPL",
            "split": "test/full",
            "rows": 5153,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/lambada_openai/test.parquet",
            "status": f"六模型×五配置运行中：{progress['task_counts']['lambada_full_ppl']}/30 完成",
            "protocol": "全文 token PPL（非末词 accuracy）；无分隔拼接；2048 non-overlap；单列披露",
            "priority": "运行中",
        },
        {
            "name": "IMDB",
            "role": "域外 PPL 压力测试",
            "split": "test/full",
            "rows": 25000,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/imdb/test.parquet",
            "status": f"六模型×五配置运行中：{progress['task_counts']['imdb_ppl']}/30 完成",
            "protocol": "text 全文无分隔拼接；2048 non-overlap；tail 丢弃；token-weighted NLL",
            "priority": "运行中",
        },
        {
            "name": "AG News",
            "role": "域外 PPL 压力测试",
            "split": "test/full",
            "rows": 7600,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/ag_news/test.parquet",
            "status": f"六模型×五配置运行中：{progress['task_counts']['ag_news_ppl']}/30 完成",
            "protocol": "text 全文无分隔拼接；2048 non-overlap；tail 丢弃；token-weighted NLL",
            "priority": "运行中",
        },
        {
            "name": "PIQA",
            "role": "ACC 主任务",
            "split": "validation/full",
            "rows": 1838,
            "path": WORKSPACE / "fpga_llm_project/data/piqa/validation_full.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc_norm",
            "priority": "正式结果",
        },
        {
            "name": "ARC-Easy",
            "role": "ACC 主任务",
            "split": "test/full",
            "rows": 2376,
            "path": WORKSPACE / "fpga_llm_project/data/lm_eval_standard/arc_easy_test.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc_norm",
            "priority": "正式结果",
        },
        {
            "name": "ARC-Challenge",
            "role": "ACC 主任务",
            "split": "test/full",
            "rows": 1172,
            "path": WORKSPACE / "fpga_llm_project/data/lm_eval_standard/arc_challenge_test.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc_norm；另有旧论文 validation 口径",
            "priority": "正式结果",
        },
        {
            "name": "BoolQ",
            "role": "ACC 主任务",
            "split": "validation/full",
            "rows": 3270,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/boolq/validation.parquet",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc",
            "priority": "正式结果",
        },
        {
            "name": "HellaSwag",
            "role": "ACC 主任务",
            "split": "validation/full",
            "rows": 10042,
            "path": WORKSPACE / "fpga_llm_project/data/hellaswag/hellaswag_val.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc_norm",
            "priority": "正式结果",
        },
        {
            "name": "WinoGrande",
            "role": "ACC 主任务",
            "split": "dev/full",
            "rows": 1267,
            "path": WORKSPACE / "fpga_llm_project/data/winogrande/winogrande_1.1/winogrande_1.1/dev.jsonl",
            "status": "六模型×五配置已完成",
            "protocol": "lm-eval 0.4.11 zero-shot；acc",
            "priority": "正式结果",
        },
        {
            "name": "OpenBookQA",
            "role": "ACC 论文扩展",
            "split": "validation/full",
            "rows": 500,
            "path": WORKSPACE / "fpga_llm_project/data/openbookqa/validation_full.jsonl",
            "status": "旧结果已排除；新五配置待补测",
            "protocol": "新协议完成前不写入结果页，不使用旧 gyrot-aligned 数值补空",
            "priority": "后续补全",
        },
        {
            "name": "GSM8K main",
            "role": "数学生成 ACC",
            "split": "test/full",
            "rows": 1319,
            "path": WORKSPACE / "fpga_llm_project/data/gsm8k_openai_main_raw/test.jsonl",
            "status": "旧 first-300 已排除；full 五配置待补测",
            "protocol": "最终答案 exact match；必须 full test，旧 first-300 不写入本工作簿",
            "priority": "后续补全",
        },
        {
            "name": "MMLU",
            "role": "57 学科 ACC",
            "split": "test/full",
            "rows": 14042,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/mmlu/all_test.parquet",
            "status": "数据已就绪，六模型×五配置待补测",
            "protocol": "0-shot 或固定 5-shot 必须分表；dev 只作 few-shot 示例，不计分",
            "priority": "后续补全",
        },
        {
            "name": "MATH-500",
            "role": "数学生成",
            "split": "test/full",
            "rows": 500,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/math500/test.jsonl",
            "status": "数据已就绪，待生成式评测",
            "protocol": "固定生成配置与数学答案归一化器；不得当作 likelihood ACC",
            "priority": "生成专项",
        },
        {
            "name": "IFEval",
            "role": "指令遵循",
            "split": "benchmark prompts/full",
            "rows": 541,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/ifeval/ifeval_input_data.jsonl",
            "status": "数据已就绪，待生成式评测",
            "protocol": "使用 pinned IFEval evaluator；与 ACC/PPL 分表",
            "priority": "生成专项",
        },
        {
            "name": "HumanEval+",
            "role": "代码生成",
            "split": "test/full",
            "rows": 164,
            "path": WORKSPACE / "fpga_llm_project/data/bape_eval/raw/humanevalplus/test.parquet",
            "status": "数据已就绪，待隔离沙箱评测",
            "protocol": "EvalPlus sandbox pass@1；不得仅比较生成文本",
            "priority": "代码专项",
        },
    ]
    headers = ["数据集", "用途", "Split", "行数", "文件大小(MiB)", "SHA-256", "本轮状态", "建议口径", "优先级", "本地路径"]
    ws.append([])
    ws.append(headers)
    for item in datasets:
        path = Path(item["path"])
        ws.append(
            [
                item["name"], item["role"], item["split"], item["rows"],
                path.stat().st_size / (1024 * 1024), sha256_file(path), item["status"],
                item["protocol"], item["priority"], str(path),
            ]
        )
    header_row = 4
    end_row = header_row + len(datasets)
    style_header(ws, header_row, len(headers))
    add_table(ws, header_row, end_row, len(headers), "DatasetCoverage")
    for row in range(header_row + 1, end_row + 1):
        ws.cell(row, 4).number_format = "#,##0"
        ws.cell(row, 5).number_format = "0.00"
        for col in range(1, len(headers) + 1):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 62
    ws.freeze_panes = "A5"
    set_widths(ws, {1: 30, 2: 20, 3: 32, 4: 12, 5: 16, 6: 68, 7: 30, 8: 65, 9: 12, 10: 90})


def notes_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("阅读说明")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "如何解读这份工作簿"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:F1")
    notes = [
        ("1. 表格结构", "“五配置总表”和“五配置逐任务”只读取本轮 GPU 正式实测；“评测口径”冻结所有可复现参数；“实验审计”记录版本、checkpoint 指纹、代码/结果 hash 和生成时完成度。"),
        ("2. 五种配置", "Full precision (BF16)、A8W8 / FP32 scale、A8W4 / FP32 scale、A8W4 / FP16 scale、A8W4 / ScaleCode。"),
        ("3. 历史数据边界", "旧论文表及 GSM8K first-300 结果不属于本轮重跑，已从交付工作簿移除。OpenBookQA、GSM8K-full、MMLU 等只有在新协议下完成五配置 GPU 实测后才能加入结果页。"),
        ("4. ACC 新正式口径", "“五配置逐任务”使用 lm-eval 0.4.11、zero-shot、完整 split；PIQA/ARC-Easy/ARC-Challenge/HellaSwag 使用 acc_norm，BoolQ/WinoGrande 使用 acc；Macro ACC 为六任务等权平均。"),
        ("5. PPL 正式口径", "WikiText-2 raw-v1 full test；原生 tokenizer；2048-token non-overlap；丢弃不足 2048 的 tail；按 scoring token 加权 NLL 后取 exp。"),
        ("6. 配置间可比性", "同一模型的五种配置复用相同 checkpoint 指纹、prompt、数据、tokenizer、batch-independent 评分逻辑与 seed，仅改变量化/scale payload。"),
        ("7. ScaleCode 判断", "A8W4 / ScaleCode 应直接与 A8W4 / FP32 scale 对比，以隔离 scale 编码误差；A8W8→A8W4 用于观察 MLP 权重 8bit→4bit 的贡献。"),
        ("8. 硬件边界", "本工作簿是软件算法 QDQ，不包含 S24、整数 Hadamard、定点 accumulator、KV12、定点 Softmax 或 FPGA bit-exact 全链路。"),
        ("9. 正式实验清单", str(RESULT_ROOT / "experiment_manifest.json")),
        ("10. 正式五配置结果", str(RESULT_ROOT / "six_model_summary.csv")),
        ("11. 五配置 Delta", str(RESULT_ROOT / "five_config_delta.csv")),
    ]
    start = 3
    for index, (title, body) in enumerate(notes, start=start):
        ws.cell(index, 1, title)
        ws.cell(index, 1).font = Font(bold=True, color=NAVY)
        ws.cell(index, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(index, 2, body)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[index].height = 54 if index < 8 else 32
    set_widths(ws, {1: 24, 2: 34, 3: 20, 4: 20, 5: 20, 6: 20})


def verify_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    expected = ["五配置总表", "ScaleCode结论", "五配置逐任务", "评测口径", "实验审计", "数据集覆盖", "阅读说明"]
    if workbook.sheetnames != expected:
        raise ValueError(f"sheet mismatch: {workbook.sheetnames}")
    if workbook["五配置总表"].max_row != 34:
        raise ValueError("formal five-config summary sheet is incomplete")
    if workbook["五配置逐任务"].max_row != 34:
        raise ValueError("formal five-config task sheet is incomplete")
    if workbook["评测口径"].max_row < 45:
        raise ValueError("evaluation protocol sheet is incomplete")
    if workbook["实验审计"].max_row < 40:
        raise ValueError("experiment audit sheet is incomplete")
    protocol_values = {
        workbook["评测口径"].cell(row, 2).value
        for row in range(1, workbook["评测口径"].max_row + 1)
    }
    for required in ("literature_2048", "20260913", "PPL = exp(sum(NLL) / scoring_tokens)"):
        if required not in protocol_values:
            raise ValueError(f"evaluation protocol missing required value: {required}")
    configs = {
        workbook["五配置总表"].cell(row, 2).value
        for row in range(5, 35)
    }
    if configs != set(FORMAL_CONFIG_DISPLAY.values()):
        raise ValueError(f"five-config labels mismatch: {sorted(configs)}")
    for model in MODELS:
        count = sum(
            workbook["五配置总表"].cell(row, 1).value == model
            for row in range(5, 35)
        )
        if count != 5:
            raise ValueError(f"{model} has {count} rows instead of five")
    workbook.close()


def main() -> None:
    args = parse_args()
    summary = read_csv(RESULT_ROOT / "six_model_summary.csv")
    deltas = read_csv(RESULT_ROOT / "six_model_delta.csv")
    if len(summary) != 30 or len(deltas) != 6:
        raise ValueError("authoritative six-model five-config matrix is incomplete")
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.title = "ScaleLLM Five-Configuration ACC/PPL Analysis"
    wb.properties.subject = "Six-model, five-configuration formal quality evaluation"
    wb.properties.creator = "ScaleLLM evaluation pipeline"
    formal_five_config_sheet(wb, summary)
    main_sheet(wb, summary, deltas)
    task_sheet(wb, summary)
    progress = extended_ppl_progress()
    protocol_sheet(wb, progress)
    audit_sheet(wb, progress)
    dataset_sheet(wb, progress)
    notes_sheet(wb)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    verify_workbook(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
