#!/usr/bin/env python3
"""Create the paper-facing workbook for the format-v3 validation study."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


WORKSPACE = Path("/home/user/EDA-202/zhuowenjun")
DEFAULT_ROOT = (
    WORKSPACE
    / "ScaleLLM"
    / "0_Software"
    / "fix_experiment"
    / "results"
    / "llama31_rope_softmax_fixedpoint_hardware128_v3"
)

NAVY = "17365D"
BLUE = "5B9BD5"
LIGHT_BLUE = "DDEBF7"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
YELLOW = "FFD966"
LIGHT_YELLOW = "FFF2CC"
RED = "C00000"
LIGHT_RED = "F4CCCC"
WHITE = "FFFFFF"
GRAY = "E7E6E6"
THIN_GRAY = Side(style="thin", color="D9E1F2")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--output", type=Path)
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def read_json(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def scalar(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, (bool, int, float)):
        return value
    if value in {"True", "False"}:
        return value == "True"
    try:
        return int(value)
    except ValueError:
        pass
    try:
        return float(value)
    except ValueError:
        return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def style_title(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(vertical="center")


def style_header(row: Any) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)


def add_table(ws: Any, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    reference = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def fit_columns(ws: Any, maximum: int = 55) -> None:
    for column in range(1, ws.max_column + 1):
        values = [
            str(ws.cell(row=row, column=column).value or "")
            for row in range(1, min(ws.max_row, 300) + 1)
        ]
        width = min(maximum, max(10, max(map(len, values), default=10) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width


def append_rows(ws: Any, rows: list[dict[str, str]], columns: list[str]) -> None:
    ws.append(columns)
    style_header(ws[1])
    for raw in rows:
        ws.append([scalar(raw.get(column)) for column in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def build_cover(
    ws: Any,
    root: Path,
    summaries: list[dict[str, str]],
    width_findings: list[dict[str, str]],
    contract: dict[str, Any] | None,
    regression: dict[str, Any] | None,
    pause: dict[str, Any] | None,
    migration: dict[str, Any] | None,
) -> None:
    ws.title = "README"
    ws.merge_cells("A1:F1")
    ws["A1"] = "ScaleLLM Llama-3.1-8B-Instruct format-v3 验证"
    style_title(ws["A1"])
    ws.row_dimensions[1].height = 28
    entries = [
        ("生成时间", datetime.now(timezone.utc).isoformat()),
        ("结果根目录", str(root)),
        ("模型", "Llama-3.1-8B-Instruct"),
        ("硬件协议", "hardware_128; context=128"),
        ("ACC", "lm-evaluation-harness 0.4.11; zero-shot continuation likelihood"),
        ("PPL", "WikiText-2 raw-v1 full test; 128-token non-overlap blocks"),
        ("Attention format", "0x48384103; total_shift [26,47]"),
        ("MLP format", "0x48384603; S37/S44/S46; total_shift [20,44]"),
        ("当前 Down 边界", "S24 Q8.16 (Down-S32 仅诊断，不是当前硬件)"),
        (
            "执行状态",
            (
                f"{pause.get('status')}; {pause.get('paused_at')}; "
                "MLP/Full ACC partial requests are not reportable ACC"
                if pause
                else "active or no explicit pause record"
            ),
        ),
        (
            "格式审计",
            contract.get("status") if contract else "missing format_v3_contract_audit.json",
        ),
        (
            "回归证据",
            regression.get("status") if regression else "missing regression_evidence.json",
        ),
        (
            "Down S32迁移审计",
            (
                f"{migration.get('status')}; "
                f"{migration.get('checks_passed')}/{migration.get('checks_total')}"
                if migration
                else "missing mlp_down_s32_migration_audit.json"
            ),
        ),
    ]
    for row_index, (label, value) in enumerate(entries, start=3):
        ws.cell(row_index, 1, label).font = Font(bold=True, color=NAVY)
        ws.cell(row_index, 2, value)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=6)
        ws.cell(row_index, 2).alignment = Alignment(wrap_text=True)
    start = len(entries) + 5
    ws.cell(start, 1, "Case 完成/判定状态")
    ws.cell(start, 1).font = Font(bold=True, color=NAVY, size=13)
    columns = (
        "case_label",
        "acc_tasks_complete",
        "macro_accuracy",
        "ppl",
        "hardware_contract_status",
        "overall",
    )
    for offset, column in enumerate(columns, start=1):
        ws.cell(start + 1, offset, column)
    style_header(ws[start + 1])
    for raw in summaries:
        ws.append([scalar(raw.get(column)) for column in columns])
    note_row = ws.max_row + 2
    ws.cell(note_row, 1, "解释边界")
    ws.cell(note_row, 1).font = Font(bold=True, color=RED)
    notes = [
        "FAIL_CONTRACT：软件按 RTL 的非法 shift 置零行为完成推理，但当前格式合同不覆盖真实激活。",
        "FAIL_PPL：即使 ACC 可接受，也不能声称语言建模质量保持。",
        "Down-S32 是因果诊断；不得作为当前 S24 硬件结果发表。",
        "只有 ACC、PPL、格式/溢出合同同时通过，才能声称当前定点链路保持精度。",
    ]
    for index, note in enumerate(notes, start=note_row + 1):
        ws.cell(index, 1, f"• {note}")
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=6)
        ws.cell(index, 1).alignment = Alignment(wrap_text=True)
    issue_row = ws.max_row + 2
    ws.cell(issue_row, 1, "尚未关闭的位宽/格式项")
    ws.cell(issue_row, 1).font = Font(bold=True, color=RED, size=13)
    open_findings = [row for row in width_findings if row.get("open") == "True"]
    if not open_findings:
        ws.cell(issue_row + 1, 1, "无")
    for offset, finding in enumerate(open_findings, start=1):
        ws.cell(
            issue_row + offset,
            1,
            f"• {finding.get('finding_id')}: {finding.get('status')} — "
            f"{finding.get('required_action')}",
        )
        ws.merge_cells(
            start_row=issue_row + offset,
            start_column=1,
            end_row=issue_row + offset,
            end_column=6,
        )
        ws.cell(issue_row + offset, 1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 27
    for column in "BCDEF":
        ws.column_dimensions[column].width = 22


def add_summary_chart(ws: Any) -> None:
    if ws.max_row < 3:
        return
    headers = {cell.value: cell.column for cell in ws[1]}
    macro_column = headers.get("macro_accuracy")
    ppl_column = headers.get("ppl")
    label_column = headers.get("case_label")
    if not macro_column or not ppl_column or not label_column:
        return
    macro_chart = BarChart()
    macro_chart.type = "col"
    macro_chart.style = 10
    macro_chart.title = "六任务 Macro ACC"
    macro_chart.y_axis.title = "Accuracy"
    macro_chart.add_data(
        Reference(ws, min_col=macro_column, min_row=1, max_row=ws.max_row),
        titles_from_data=True,
    )
    macro_chart.set_categories(
        Reference(ws, min_col=label_column, min_row=2, max_row=ws.max_row)
    )
    macro_chart.height = 7
    macro_chart.width = 12
    ws.add_chart(macro_chart, f"W2")

    ppl_chart = BarChart()
    ppl_chart.type = "col"
    ppl_chart.style = 11
    ppl_chart.title = "WikiText-2 PPL"
    ppl_chart.y_axis.title = "Perplexity"
    ppl_chart.add_data(
        Reference(ws, min_col=ppl_column, min_row=1, max_row=ws.max_row),
        titles_from_data=True,
    )
    ppl_chart.set_categories(
        Reference(ws, min_col=label_column, min_row=2, max_row=ws.max_row)
    )
    ppl_chart.height = 7
    ppl_chart.width = 12
    ws.add_chart(ppl_chart, f"W17")


def main() -> None:
    args = parse_args()
    root = args.root.resolve()
    output = (
        args.output.resolve()
        if args.output
        else root / "ScaleLLM_Llama31_format_v3_validation.xlsx"
    )
    inputs = {
        "case_summary": root / "case_summary.csv",
        "case_metrics": root / "case_metrics.csv",
        "trace_coverage": root / "trace_coverage.csv",
        "node_trace": root / "node_trace.csv",
        "hardware_audit": root / "hardware_audit.csv",
        "linear_audit": root / "linear_audit_by_module.csv",
        "format_contract": root / "format_v3_contract_audit.csv",
        "width_findings": root / "width_findings.csv",
    }
    summaries = read_csv(inputs["case_summary"])
    metrics = read_csv(inputs["case_metrics"])
    width_findings = read_csv(inputs["width_findings"])
    contract_json_path = root / "format_v3_contract_audit.json"
    contract = read_json(contract_json_path) if contract_json_path.is_file() else None
    regression_json_path = root / "regression_evidence.json"
    regression = read_json(regression_json_path) if regression_json_path.is_file() else {}
    pause_json_path = root / "validation_pause.json"
    pause = read_json(pause_json_path) if pause_json_path.is_file() else None
    migration_json_path = root / "mlp_down_s32_migration_audit.json"
    migration = (
        read_json(migration_json_path) if migration_json_path.is_file() else None
    )
    pause_rows = []
    if pause:
        for worker in pause.get("workers", []):
            pause_rows.append(
                {
                    "status": pause.get("status"),
                    "paused_at": pause.get("paused_at"),
                    "reason": pause.get("reason"),
                    **worker,
                    "interpretation": pause.get("interpretation"),
                }
            )

    workbook = Workbook()
    build_cover(
        workbook.active,
        root,
        summaries,
        width_findings,
        contract,
        regression,
        pause,
        migration,
    )

    sheet_specs = [
        (
            "Run Status",
            pause_rows,
            [
                "status",
                "paused_at",
                "reason",
                "case",
                "task",
                "likelihood_requests_processed",
                "likelihood_requests_total",
                "progress_percent",
                "manifest_present",
                "valid_accuracy_available",
                "log",
                "log_sha256",
                "interpretation",
            ],
        ),
        (
            "MLP S32 Migration",
            (migration or {}).get("checks", []),
            [
                "check_id",
                "component",
                "pass",
                "severity",
                "expected",
                "observed",
                "required_action",
                "sources",
                "source_sha256",
            ],
        ),
        (
            "Case Summary",
            summaries,
            [
                "case",
                "case_label",
                "acc_tasks_complete",
                "acc_tasks_expected",
                "macro_accuracy",
                "baseline_macro_accuracy",
                "macro_delta_pp",
                "acc_limit_pp",
                "acc_quality_pass",
                "ppl",
                "baseline_ppl",
                "ppl_delta_percent",
                "ppl_limit_percent",
                "ppl_quality_pass",
                "hardware_contract_status",
                "format_errors",
                "guard_overflow",
                "accumulator_overflow",
                "linear_output_saturation",
                "overall",
                "manifest",
            ],
        ),
        (
            "ACC",
            [row for row in metrics if row.get("kind") == "acc"],
            [
                "case",
                "case_label",
                "task",
                "primary_metric",
                "value",
                "baseline_value",
                "delta_pp",
                "stderr",
                "samples",
                "status",
                "hardware_contract_status",
                "manifest",
            ],
        ),
        (
            "PPL",
            [row for row in metrics if row.get("kind") == "ppl"],
            [
                "case",
                "case_label",
                "value",
                "baseline_value",
                "absolute_delta",
                "relative_delta_percent",
                "evaluated_blocks",
                "scoring_tokens",
                "full_stream",
                "status",
                "hardware_contract_status",
                "manifest",
            ],
        ),
        (
            "Trace Coverage",
            read_csv(inputs["trace_coverage"]),
            [
                "case",
                "trace_domain",
                "required_layers",
                "observed_layers",
                "required_nodes_per_layer",
                "complete",
                "manifest",
            ],
        ),
        (
            "Node Trace",
            read_csv(inputs["node_trace"]),
            [
                "case",
                "trace_domain",
                "layer",
                "trace_call",
                "tokens",
                "node",
                "elements",
                "cosine",
                "nmse",
                "mae",
                "max_abs_error",
                "reference_abs_max",
                "valid_scores",
                "reference_score_boundary_count",
                "manifest",
            ],
        ),
        (
            "Hardware Audit",
            read_csv(inputs["hardware_audit"]),
            ["case", "metric", "value", "manifest"],
        ),
        (
            "Linear Audit",
            read_csv(inputs["linear_audit"]),
            [
                "case",
                "module",
                "linear_total_shift_min",
                "linear_total_shift_max",
                "linear_group_guard_overflow_s37",
                "linear_software_group_guard_overflow",
                "linear_accumulator_overflow",
                "linear_software_accumulator_overflow",
                "linear_output_saturation",
                "manifest",
            ],
        ),
        (
            "Format Contract",
            read_csv(inputs["format_contract"]),
            [
                "check_id",
                "component",
                "severity",
                "expected",
                "observed",
                "pass",
                "source",
            ],
        ),
        (
            "Width Findings",
            width_findings,
            [
                "finding_id",
                "component",
                "active_format",
                "status",
                "severity",
                "open",
                "evidence",
                "required_action",
                "source",
            ],
        ),
        (
            "Regression",
            regression.get("suites", []),
            [
                "suite",
                "status",
                "expected_tests",
                "observed_tests",
                "returncode",
                "output",
            ],
        ),
    ]
    for index, (name, rows, columns) in enumerate(sheet_specs, start=1):
        ws = workbook.create_sheet(name)
        append_rows(ws, rows, columns)
        add_table(ws, f"ValidationTable{index}")
        fit_columns(ws)
        if name == "Case Summary":
            add_summary_chart(ws)
            status_column = columns.index("overall") + 1
            status_range = (
                f"{get_column_letter(status_column)}2:"
                f"{get_column_letter(status_column)}{max(2, ws.max_row)}"
            )
            ws.conditional_formatting.add(
                status_range,
                CellIsRule(
                    operator="equal",
                    formula=['"PASS"'],
                    fill=PatternFill("solid", fgColor=LIGHT_GREEN),
                ),
            )
            for failure_status in (
                "FAIL_CONTRACT",
                "FAIL_ACC_PPL",
                "FAIL_ACC",
                "FAIL_PPL",
            ):
                ws.conditional_formatting.add(
                    status_range,
                    CellIsRule(
                        operator="equal",
                        formula=[f'"{failure_status}"'],
                        fill=PatternFill("solid", fgColor=LIGHT_RED),
                    ),
                )
        if name == "Node Trace" and ws.max_row >= 2:
            cosine_column = columns.index("cosine") + 1
            cell_range = (
                f"{get_column_letter(cosine_column)}2:"
                f"{get_column_letter(cosine_column)}{ws.max_row}"
            )
            ws.conditional_formatting.add(
                cell_range,
                ColorScaleRule(
                    start_type="num",
                    start_value=0.8,
                    start_color=LIGHT_RED,
                    mid_type="num",
                    mid_value=0.99,
                    mid_color=LIGHT_YELLOW,
                    end_type="num",
                    end_value=1.0,
                    end_color=LIGHT_GREEN,
                ),
            )

    provenance = workbook.create_sheet("Provenance")
    provenance.append(["artifact", "path", "exists", "sha256"])
    style_header(provenance[1])
    provenance_paths = list(inputs.values()) + [
        contract_json_path,
        regression_json_path,
        root / "precision_validation_report.md",
        root.parent
        / "six_model_algorithm_literature_2048"
        / "model_identities"
        / "Llama-3.1-8B-Instruct.json",
    ]
    for path in provenance_paths:
        provenance.append(
            [
                path.name,
                str(path),
                path.is_file(),
                sha256_file(path) if path.is_file() else None,
            ]
        )
    add_table(provenance, "ValidationProvenance")
    fit_columns(provenance)
    provenance.freeze_panes = "A2"

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output)

    # Reopen the final ZIP container so a truncated/corrupt workbook cannot be
    # reported as complete merely because save() returned.
    verified = load_workbook(output, read_only=True, data_only=False)
    sheet_rows = {sheet: verified[sheet].max_row for sheet in verified.sheetnames}
    required = {
        "README",
        "Case Summary",
        "ACC",
        "PPL",
        "Trace Coverage",
        "Node Trace",
        "Hardware Audit",
        "Linear Audit",
        "Format Contract",
        "Width Findings",
        "Regression",
        "Provenance",
    }
    if not required.issubset(verified.sheetnames):
        raise RuntimeError(f"workbook sheets missing: {sorted(required - set(verified.sheetnames))}")
    verified.close()
    manifest = {
        "schema": "scalellm_fixedpoint_validation_workbook_v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "status": "complete",
        "workbook": str(output),
        "size": output.stat().st_size,
        "sha256": sha256_file(output),
        "sheet_rows": sheet_rows,
    }
    manifest_path = root / "validation_workbook_manifest.json"
    temporary_manifest = manifest_path.with_suffix(manifest_path.suffix + ".tmp")
    temporary_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary_manifest.replace(manifest_path)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
