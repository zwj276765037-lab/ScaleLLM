#!/usr/bin/env python3
"""Build paired task-quality tables from read-only historical prediction records."""
from pathlib import Path
from collections import Counter, OrderedDict
import csv
import hashlib
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OLD = Path("/home/user/EDA-202/zhuowenjun/ScaleLLM/0_Software/results/v5_fullchain/Llama-3.1-8B-Instruct")
FINAL_TABLES = OLD.parents[1] / "Llama-3.1-8B-Instruct"
SOURCE_FILES = [OLD / arm / name for arm in ["bf16", "v5"]
                for name in ["mcqa_records.jsonl", "manifest.json"]]
FIELDS = [
    "Task", "Paired_N", "Reference_correct", "Candidate_correct", "Both_correct",
    "Correct_to_wrong", "Wrong_to_correct", "Both_wrong",
    "Correct_answer_retention_percent", "Evidence_status",
]
LABELS = {"piqa": "PIQA", "arc_challenge": "ARC-Challenge", "openbookqa": "OpenBookQA",
          "hellaswag": "HellaSwag", "winogrande": "WinoGrande"}


def identity(path):
    return {"path": str(path.resolve()), "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def read_records(arm):
    records = {}
    with (OLD / arm / "mcqa_records.jsonl").open() as stream:
        for line in stream:
            r = json.loads(line)
            key = (r["task"], r["sample_id"])
            assert key not in records, f"Duplicate prediction key: {key}"
            assert int(r["active_correct"]) == int(r["active_pred"] == r["gold_index"])
            records[key] = r
    return records


def choices(record):
    c = record["choice_scores_json"]
    c = json.loads(c) if isinstance(c, str) else c
    return [(v["suffix"], v["num_completion_tokens"]) for v in c]


def write_csv(name, fields, rows):
    with (HERE / name).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def styled_sheet(workbook, title, rows, widths):
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234B6B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 44
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5F9")
        sheet.row_dimensions[row[0].row].height = 34
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    return sheet


def main():
    sources = [identity(p) for p in SOURCE_FILES]
    final_sources = []
    for arm in ["bf16", "scalecode"]:
        path = FINAL_TABLES / f"accuracy_{arm}_fullsplit.csv"
        sources.append(identity(path))
        with path.open(newline="") as stream:
            for r in csv.DictReader(stream):
                target = Path(r["records_artifact"])
                final_sources.append({
                    "arm": arm, "task": r["task"], "declared_samples": int(r["samples"]),
                    "summary_csv": str(path), "records_artifact": str(target),
                    "records_file_exists": target.is_file(),
                })
    (HERE / "final_acc_source_availability.json").write_text(
        json.dumps(final_sources, ensure_ascii=False, indent=2) + "\n")
    reference, candidate = read_records("bf16"), read_records("v5")
    assert reference.keys() == candidate.keys()
    counts = OrderedDict()
    for key, a in reference.items():
        b = candidate[key]
        assert a["sample_index"] == b["sample_index"] and a["gold_index"] == b["gold_index"]
        assert choices(a) == choices(b), f"Candidate identity differs at {key}"
        bucket = counts.setdefault(key[0], Counter())
        ac, bc = bool(a["active_correct"]), bool(b["active_correct"])
        bucket["Paired_N"] += 1
        bucket["Reference_correct"] += ac
        bucket["Candidate_correct"] += bc
        bucket["Both_correct"] += ac and bc
        bucket["Correct_to_wrong"] += ac and not bc
        bucket["Wrong_to_correct"] += not ac and bc
        bucket["Both_wrong"] += not ac and not bc
    total = sum(counts.values(), Counter())
    assert total["Paired_N"] == 1499
    manifests = {arm: json.loads((OLD / arm / "manifest.json").read_text())
                 for arm in ["bf16", "v5"]}
    assert manifests["v5"]["status"] == "complete_invalid_format_error"
    output = []
    for task, c in list(counts.items()) + [("TOTAL", total)]:
        assert c["Both_correct"] + c["Correct_to_wrong"] == c["Reference_correct"]
        assert c["Both_correct"] + c["Wrong_to_correct"] == c["Candidate_correct"]
        assert sum(c[k] for k in ["Both_correct", "Correct_to_wrong",
                                 "Wrong_to_correct", "Both_wrong"]) == c["Paired_N"]
        output.append({
            "Task": LABELS.get(task, task), **dict(c),
            "Correct_answer_retention_percent":
                100 * c["Both_correct"] / c["Reference_correct"] if c["Reference_correct"] else None,
            "Evidence_status": "HISTORICAL_DIAGNOSTIC_ONLY_INVALID_FORMAT",
        })
    write_csv("historical_paired_answer_quality.csv", FIELDS, output)
    pending = {field: None for field in FIELDS}
    pending.update({
        "Task": "Current multistage-rounding comparison",
        "Evidence_status": "MISSING_MATCHED_PER_QUESTION_RESULTS",
    })
    write_csv("current_rounding_answer_quality.csv", FIELDS, [pending])
    record_fields = ["task", "sample_id", "sample_index", "gold_index",
                     "reference_prediction", "candidate_prediction",
                     "reference_correct", "candidate_correct", "transition"]
    paired = []
    for key, a in reference.items():
        b = candidate[key]
        ac, bc = int(a["active_correct"]), int(b["active_correct"])
        paired.append({
            "task": a["task"], "sample_id": a["sample_id"], "sample_index": a["sample_index"],
            "gold_index": a["gold_index"], "reference_prediction": a["active_pred"],
            "candidate_prediction": b["active_pred"],
            "reference_correct": ac, "candidate_correct": bc,
            "transition": f"{'correct' if ac else 'wrong'}_to_{'correct' if bc else 'wrong'}",
        })
    write_csv("historical_paired_question_records.csv", record_fields, paired)
    notes = [
        ["项目", "说明"],
        ["数据有效性", "历史V5运行被标记为 complete_invalid_format_error；Historical_pairs仅作诊断，不能用于证明当前舍入方法保持输出质量。"],
        ["实际数据范围", "Llama-3.1-8B-Instruct，1499道选择题，BF16与旧V5；候选似然评分，不是自由生成长答案。"],
        ["配对检查", "task/sample_id唯一且两臂集合相同；sample_index、gold_index、候选文本和候选token数一致；未丢弃不利结果。"],
        ["正确答案保持率", "Both_correct / Reference_correct × 100%；分母只包含参考版答对的题，不是所有题，也不是候选ACC。"],
        ["原数据异常", "V5累计linear_total_shift_format_error=517、linear_output_saturation=34239；全部原始记录保留。"],
        ["协议溯源缺口", "V5合并来源缺少source0_manifest.json，前三任务原始args/model不能由该来源文件完整闭合；完整prompt/input_ids及数据集、checkpoint、代码SHA未保存。"],
        ["当前方法", "Current_rounding单独预留“相同量化与LUT规则、仅改变中间舍入”的配对结果；当前没有符合这一条件的逐题数据。"],
        ["最终ACC来源", "旧目录最终ACC CSV另引用3份全量gyrot_aligned_mcqa_records.jsonl；本次检查原路径均不存在。见final_acc_source_availability.json。"],
        ["N/A", "没有数据或尚未验证，不等于0。历史配对统计不能移填到Current_rounding。"],
        ["任务与评分", "每任务前最多300题，ARC-Challenge为299；前四任务score_char_norm，WinoGrande为score_sum。不是完整测试集统计。"],
        ["新实验", "此次只读取旧数据并计算配对表，没有启动模型/GPU或修改旧工程。"],
        ["文件", "CSV、逐题配对明细和source_manifest.json保存在本工作簿旁；本表来源不等同于M005数值实验。"],
    ]
    wb = Workbook()
    wb.remove(wb.active)
    styled_sheet(wb, "Read_me", notes, [24, 116])
    rows = [FIELDS] + [[row[f] if row[f] is not None else "N/A" for f in FIELDS] for row in output]
    sheet = styled_sheet(wb, "Historical_pairs", rows, [20, 13, 18, 18, 16, 18, 18, 16, 24, 47])
    for row in sheet.iter_rows(min_row=2):
        row[8].number_format = '0.00"%"'
        row[9].font = Font(color="A12722")
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True, color="234B6B")
    styled_sheet(wb, "Current_rounding", [FIELDS, [pending[f] if pending[f] is not None else "N/A" for f in FIELDS]],
                 [38, 13, 18, 18, 16, 18, 18, 16, 24, 47])
    styled_sheet(wb, "Question_pairs",
                 [record_fields] + [[r[f] for f in record_fields] for r in paired],
                 [18, 30, 16, 15, 22, 22, 20, 20, 25])
    path = HERE / "输出质量保持数据表.xlsx"
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert check["Historical_pairs"].max_row == 7 and check["Question_pairs"].max_row == 1500
    assert check["Current_rounding"]["B2"].value == "N/A"
    check.close()
    md = [
        "# 输出质量保持数据表",
        "",
        "**这是已找到的旧版诊断数据，不是当前舍入方法的合格结果。** "
        "BF16与旧V5共有1499道可配对选择题。V5运行状态为`complete_invalid_format_error`，"
        "存在517次移位格式错误和34239次线性输出饱和；不能据此支持当前方法的论文结论。",
        "",
        "[Excel工作簿](输出质量保持数据表.xlsx) · [配对汇总CSV](historical_paired_answer_quality.csv)"
        " · [1499题配对明细](historical_paired_question_records.csv)",
        "",
        "| 任务 | 配对题数 | BF16正确 | 旧V5正确 | 两版均正确 | 对→错 | 错→对 | 两版均错 | 正确答案保持率 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for r in output:
        md.append(f"| {r['Task']} | {r['Paired_N']} | {r['Reference_correct']} | "
                  f"{r['Candidate_correct']} | {r['Both_correct']} | {r['Correct_to_wrong']} | "
                  f"{r['Wrong_to_correct']} | {r['Both_wrong']} | {r['Correct_answer_retention_percent']:.2f}% |")
    md += [
        "",
        "保持率 = 两版均正确 / BF16正确；“两版均错”不表示两版选择了相同错误答案。"
        "TOTAL为题目合池统计，不是任务等权平均。",
        "",
        "题目ID、题目索引、金标、候选文本和候选token数均一致；"
        "正确性由各自active_pred与gold_index重算。此处为候选选项评分，不是自由生成回答质量评测。"
        "前四任务采用score_char_norm，WinoGrande采用score_sum；每任务前最多300题，ARC-Challenge为299题。",
        "",
        "V5合并来源还缺少source0_manifest.json，前三任务的原始args/model没有从该文件完整闭合；"
        "完整prompt/input_ids及数据集、checkpoint、代码SHA未保存。"
        "即使配对计数可复算，格式错误和来源缺口仍使该表只能作为历史诊断。",
        "",
        "## 当前多级舍入对照",
        "",
        "论文目标对照应固定量化权重、ScaleCode、LUT、输入和评分协议，"
        "只比较保留中间精度参照与当前多级舍入实现。"
        "现有103,700个数值预测位置没有逐题金标和任务判分，不能替代这里的题目样本。",
        "旧目录中的最终ACC汇总另外指向3份全量逐题文件，本次检查原路径均不存在；"
        "源CSV和具体缺失路径见[最终ACC逐题来源检查](final_acc_source_availability.json)。"
        "不能用汇总正确数反推出两版均正确、对→错或错→对。",
        "",
        "| 配对题数 | 参考正确 | 两版均正确 | 对→错 | 错→对 | 正确答案保持率 | 状态 |",
        "|---:|---:|---:|---:|---:|---:|---|",
        "| N/A | N/A | N/A | N/A | N/A | N/A | 缺少符合对照条件的逐题结果 |",
        "",
        "[当前对照待补表](current_rounding_answer_quality.csv)。N/A表示缺失，不表示0。"
        "历史BF16/旧V5结果不填入该对照。此次未新跑模型。",
        "",
        "来源与文件SHA见[source_manifest.json](source_manifest.json)；"
        "配对汇总可通过[build_tables.py](build_tables.py)重新生成。",
        "",
    ]
    (HERE / "输出质量保持数据表.md").write_text("\n".join(md))
    for r in sources:
        assert identity(Path(r["path"]))["sha256"] == r["sha256"]
    outputs = ["输出质量保持数据表.xlsx", "输出质量保持数据表.md",
               "historical_paired_answer_quality.csv", "current_rounding_answer_quality.csv",
               "historical_paired_question_records.csv", "final_acc_source_availability.json"]
    manifest = {
        "execution_state": "completed",
        "current_rounding_quality_verdict": "not_measured",
        "historical_evidence_status": "invalid_format_diagnostic_only",
        "sources": sources, "builder": identity(Path(__file__)),
        "outputs": [identity(HERE / name) for name in outputs],
        "paired_count": len(paired), "rows": output,
        "source_files_unchanged": True, "model_or_gpu_execution": False,
        "counts_verified": True, "workbook_reopened_and_verified": True,
    }
    (HERE / "source_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"workbook": str(path), "rows": output}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
